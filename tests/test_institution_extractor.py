from io import BytesIO

from docx import Document

from config.institutions import KNOWN_INSTITUTIONS
from services.document_parser import parse_document_bytes
from services.institution_extractor import extract_institution_name, institution_is_resolved


def parsed_text(text, file_name="self_study.txt", section="Title page", file_type="txt"):
    return [{
        "file_name": file_name,
        "file_type": file_type,
        "text_blocks": [{"section": section, "text": text}],
        "tables": [],
        "detected_document_type": "self study",
        "extraction_warnings": [],
        "ocr_used": False,
    }]


def test_exact_institution_label_extraction():
    result = extract_institution_name(
        parsed_text("Name of Institution: University of Lagos"),
        KNOWN_INSTITUTIONS,
    )

    assert result.status == "confirmed"
    assert result.matched_known_institution == "University of Lagos"
    assert result.confidence >= 0.85


def test_heading_extraction():
    result = extract_institution_name(
        parsed_text("AHMADU BELLO UNIVERSITY\nFaculty of Computing", "cover.txt"),
        KNOWN_INSTITUTIONS,
    )

    assert result.status == "confirmed"
    assert result.matched_known_institution == "Ahmadu Bello University"


def test_docx_table_extraction():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name of Institution"
    table.cell(0, 1).text = "University of Ibadan"
    stream = BytesIO()
    document.save(stream)

    parsed = parse_document_bytes(stream.getvalue(), "self_study.docx", "docx")
    result = extract_institution_name([parsed], KNOWN_INSTITUTIONS)

    assert result.status == "confirmed"
    assert result.matched_known_institution == "University of Ibadan"
    assert result.source_page_or_section.startswith("Table")


def test_pdf_title_page_extraction_from_parsed_pdf_text():
    result = extract_institution_name(
        parsed_text("UNIVERSITY OF LAGOS\nSelf-Study Form", "form.pdf", "Title page", "pdf"),
        KNOWN_INSTITUTIONS,
    )

    assert result.status == "confirmed"
    assert result.source_file == "form.pdf"
    assert result.source_page_or_section == "Title page"
    assert result.matched_known_institution == "University of Lagos"


def test_fuzzy_match_punctuation_difference():
    result = extract_institution_name(
        parsed_text("Name of University: University of Nigeria, Nsukka"),
        KNOWN_INSTITUTIONS,
    )

    assert result.status == "confirmed"
    assert result.matched_known_institution == "University of Nigeria Nsukka"


def test_known_abbreviation_with_supporting_evidence_is_confirmed():
    result = extract_institution_name(
        parsed_text("Institution: UNILAG\nUniversity of Lagos Self-Study Form"),
        KNOWN_INSTITUTIONS,
    )

    assert result.status == "confirmed"
    assert result.matched_known_institution == "University of Lagos"


def test_ambiguous_abbreviation_does_not_auto_confirm():
    result = extract_institution_name(
        parsed_text("Institution: UI"),
        KNOWN_INSTITUTIONS,
    )

    assert result.status == "needs_review"
    assert result.matched_known_institution == "University of Ibadan"
    assert result.confidence < 0.85


def test_missing_institution():
    result = extract_institution_name(
        parsed_text("Programme staffing and curriculum evidence only."),
        KNOWN_INSTITUTIONS,
    )

    assert result.status == "missing"
    assert result.matched_known_institution == ""


def test_conflicting_documents():
    result = extract_institution_name(
        parsed_text("Name of Institution: University of Lagos", "one.txt")
        + parsed_text("Name of Institution: University of Ibadan", "two.txt"),
        KNOWN_INSTITUTIONS,
    )

    assert result.status == "conflicting"
    assert len(result.alternative_candidates) >= 2


def test_user_correction_preserves_original_evidence():
    result = extract_institution_name(
        parsed_text("Institution: UNILAG"),
        KNOWN_INSTITUTIONS,
    )
    evidence = result.to_field_evidence(user_confirmed=True)
    evidence.update({
        "confirmed_value": "University of Lagos",
        "corrected_by_user": True,
        "correction_timestamp": "2026-07-13T10:00:00",
    })

    assert evidence["extracted_value"] == "University of Lagos"
    assert evidence["confirmed_value"] == "University of Lagos"
    assert evidence["corrected_by_user"] is True
    assert evidence["evidence_text"] == "Institution: UNILAG"


def test_prediction_is_blocked_until_institution_is_resolved():
    assert institution_is_resolved("") is False
    assert institution_is_resolved(None) is False
    assert institution_is_resolved("University of Lagos") is True
